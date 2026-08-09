"""Generate the compact 2-3 page Project Update 1 DOCX report."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports" / "Project_Update_1_Report.docx"

# Standard-business-brief preset with a named course_page_limit override:
# 0.55-inch margins, Arial 9 pt, and compressed paragraph rhythm are used to
# satisfy the instructor's explicit 2-3 page limit while retaining readability.
INK = RGBColor(28, 45, 66)
BLUE = RGBColor(42, 91, 134)
MUTED = RGBColor(90, 100, 112)
LIGHT = "D9E7F2"


PAPERS = [
    {
        "citation": "1. Oh et al. (2018) - CNN-LSTM with variable-length heartbeats",
        "summary": (
            "Oh et al. combined convolutional layers with LSTM units to classify variable-length ECG "
            "segments from MIT-BIH into normal rhythm, LBBB, RBBB, atrial premature beats, and premature "
            "ventricular contractions. CNN layers learned local waveform morphology, while the LSTM modeled "
            "sequential variation without forcing every beat to the same duration. Ten-fold cross-validation "
            "produced 98.10% accuracy, 97.50% sensitivity, and 98.70% specificity. The study directly supports "
            "our hybrid architecture. However, its five selected labels do not follow the complete AAMI "
            "superclass scheme, and random folds may contain beats from the same patient. The resulting gap "
            "is a leakage-aware, record-wise CNN-LSTM evaluation with explicit class-imbalance reporting."
        ),
    },
    {
        "citation": "2. Kachuee, Fazeli, and Sarrafzadeh (2018) - Deep transferable representation",
        "summary": (
            "Kachuee et al. trained a 13-layer one-dimensional residual CNN on MIT-BIH heartbeats grouped "
            "into five AAMI classes, then transferred the learned representation to myocardial-infarction "
            "classification on the PTB database. Their pipeline used automated beat extraction, balanced "
            "augmentation, residual blocks, and a softmax classifier. Average accuracy reached 93.4% for "
            "arrhythmia and 95.9% for myocardial infarction, demonstrating that ECG features can transfer "
            "between tasks. The model does not explicitly capture longer beat-to-beat dependencies, and "
            "balanced test sampling differs from natural clinical prevalence. Our project can retain its "
            "efficient residual-style feature extraction while adding LSTM context and reporting metrics on "
            "an untouched, record-separated test distribution."
        ),
    },
    {
        "citation": "3. Mousavi, Afghah, and Acharya (2019) - Inter/intra-patient sequence-to-sequence model",
        "summary": (
            "Mousavi et al. paired a three-layer one-dimensional CNN with an encoder-decoder sequence model "
            "to label ordered MIT-BIH heartbeats. Beats were normalized, resized to 280 samples, and grouped "
            "under AAMI categories; SMOTE addressed minority classes. Importantly, the authors evaluated both "
            "intra-patient and inter-patient protocols. In the inter-patient setting, positive predictive value "
            "and sensitivity reached 92.57% and 88.94% for class S, and 99.50% and 99.94% for class V. This "
            "study shows that temporal context and patient separation matter. Its synthetic oversampling and "
            "annotation-centered segmentation may still limit deployment realism. Our work should apply "
            "balancing only to training records and disclose reliance on reference R-peaks."
        ),
    },
    {
        "citation": "4. Romdhane et al. (2020) - Deep CNN with focal loss",
        "summary": (
            "Romdhane et al. proposed a deep one-dimensional CNN and focal loss for five-class AAMI heartbeat "
            "classification on MIT-BIH and INCART. Their segmentation began at each R-peak and extended for "
            "1.2 times the local median RR interval, avoiding fixed morphology assumptions and explicit "
            "denoising. Focal loss emphasized difficult minority samples. The system reported 98.41% accuracy, "
            "98.38% F1-score, 98.37% precision, and 98.41% recall. The work is valuable for our highly "
            "imbalanced label distribution, but a single aggregate score can conceal patient-level failures, "
            "and annotation-derived R-peaks simplify real deployment. We should therefore compare weighted "
            "loss against focal loss using macro-F1, per-class sensitivity, and a record-wise held-out set."
        ),
    },
    {
        "citation": "5. Hassan et al. (2022) - CNN with bidirectional LSTM",
        "summary": (
            "Hassan et al. combined one-dimensional convolution with bidirectional LSTM layers to classify N, "
            "S, V, F, and Q heartbeats. CNN blocks captured local ECG morphology, while forward and backward "
            "recurrent passes modeled temporal dependencies. The model was tested separately on MIT-BIH and "
            "the St-Petersburg database, reaching 98% test accuracy and about 91% sensitivity on MIT-BIH, but "
            "lower performance on St-Petersburg. The external decline is important evidence that benchmark "
            "accuracy does not guarantee cross-database generalization. The paper does not fully resolve class "
            "imbalance or acquisition-domain differences. Our project should treat CNN-BiLSTM as a strong "
            "comparison model, prioritize minority-class recall, and reserve external validation as future "
            "work."
        ),
    },
    {
        "citation": "6. Yang, Liu, and Zhang (2022) - Lightweight CNN-BiLSTM with weighted loss",
        "summary": (
            "Yang et al. developed a lightweight depthwise-separable CNN followed by BiLSTM and a weighted loss "
            "for five AAMI heartbeat groups. Raw MIT-BIH signals were segmented with minimal preprocessing, "
            "and ten-fold validation assessed robustness. The model achieved 99.33% accuracy, 93.67% "
            "sensitivity, 99.18% specificity, 89.85% positive predictive value, and 91.65% F1-score. The gap "
            "between accuracy and F1 confirms that imbalance can make overall accuracy optimistic. Although "
            "the compact design is suitable for wearable devices, validation remains confined to MIT-BIH and "
            "fold construction may influence results. Our project can adopt class-weighted training and model "
            "efficiency reporting while enforcing record separation and comparing macro-F1 with accuracy."
        ),
    },
    {
        "citation": "7. Rashed-Al-Mahfuz et al. (2021) - Time-frequency VGG16 with SHAP",
        "summary": (
            "Rashed-Al-Mahfuz et al. converted temporal ECG beats into time-frequency representations, classified "
            "them with a customized VGG16 CNN, and used SHAP to identify influential frequency regions. On "
            "MIT-BIH, the reported accuracy was 100% for two-to-four classes and 99.90% for five classes; a "
            "mixed external test reported 99.91%. Interpretability is the paper's major contribution because it "
            "links predictions to frequency components rather than presenting only labels. However, scalogram "
            "or Hilbert-Huang conversion increases preprocessing and computation, and near-perfect results "
            "require careful split auditing. Our one-dimensional CNN-LSTM should remain computationally simpler, "
            "but later work can add saliency or gradient-based explanations and verify them on record-separated "
            "patients."
        ),
    },
    {
        "citation": "8. Sun et al. (2024) - CNN-LSTM with squeeze-and-excitation attention",
        "summary": (
            "Sun et al. combined ensemble empirical mode decomposition denoising, CNN feature extraction, LSTM "
            "temporal modeling, and a squeeze-and-excitation channel-attention module on MIT-BIH. The attention "
            "block reweighted feature channels, and comparisons against LSTM, CNN-LSTM, and LSTM-attention "
            "baselines attributed gains to the combined design. The reported accuracy was 98.5%, with class "
            "precision above 97%, recall above 98%, and F1 above 0.98. The architecture is directly aligned "
            "with our topic, but EEMD and attention add complexity, and the evaluation offers limited evidence "
            "of cross-database robustness. Our project should first establish simpler CNN and LSTM baselines, "
            "then add attention only if record-wise results justify the extra complexity."
        ),
    },
]


REFERENCES = [
    "Oh, S. L., Ng, E. Y. K., Tan, R. S., & Acharya, U. R. (2018). Automated diagnosis of arrhythmia using combination of CNN and LSTM techniques with variable length heart beats. Computers in Biology and Medicine, 102, 278-287. https://doi.org/10.1016/j.compbiomed.2018.06.002",
    "Kachuee, M., Fazeli, S., & Sarrafzadeh, M. (2018). ECG heartbeat classification: A deep transferable representation. IEEE ICHI, 443-444. https://doi.org/10.1109/ICHI.2018.00092",
    "Mousavi, S., Afghah, F., & Acharya, U. R. (2019). Inter- and intra-patient ECG heartbeat classification for arrhythmia detection: A sequence-to-sequence deep learning approach. ICASSP, 1308-1312. https://doi.org/10.1109/ICASSP.2019.8683140",
    "Romdhane, T. F., Alhichri, H., Ouni, R., & Atri, M. (2020). Electrocardiogram heartbeat classification based on a deep convolutional neural network and focal loss. Computers in Biology and Medicine, 123, 103866. https://doi.org/10.1016/j.compbiomed.2020.103866",
    "Hassan, S. U., Zahid, M. S. M., Abdullah, T. A. A., & Husain, K. (2022). Classification of cardiac arrhythmia using a convolutional neural network and bi-directional long short-term memory. DIGITAL HEALTH, 8. https://doi.org/10.1177/20552076221102766",
    "Yang, M., Liu, W., & Zhang, H. (2022). A robust multiple heartbeats classification with weight-based loss based on convolutional neural network and bidirectional long short-term memory. Frontiers in Physiology, 13, 982537. https://doi.org/10.3389/fphys.2022.982537",
    "Rashed-Al-Mahfuz, M., et al. (2021). Deep convolutional neural networks based ECG beats classification to diagnose cardiovascular conditions. Biomedical Engineering Letters, 11(2), 147-162. https://doi.org/10.1007/s13534-021-00185-w",
    "Sun, A., Hong, W., Li, J., & Mao, J. (2024). An arrhythmia classification model based on a CNN-LSTM-SE algorithm. Sensors, 24(19), 6306. https://doi.org/10.3390/s24196306",
]


def set_font(run, size=9, bold=False, color=None, italic=False):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    prefix = paragraph.add_run("Page ")
    set_font(prefix, size=7.5, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    rpr.append(fonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "15")
    rpr.append(sz)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    fld.append(run)
    paragraph._p.append(fld)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2A5B86")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(underline)
    run.append(rpr)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_paper_review(document, paper):
    heading = document.add_paragraph(style="Paper Heading")
    heading.paragraph_format.keep_with_next = True
    run = heading.add_run(paper["citation"])
    set_font(run, size=9.2, bold=True, color=INK)

    paragraph = document.add_paragraph(style="Paper Summary")
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(paper["summary"])
    set_font(run, size=8.8)


for paper in PAPERS:
    count = len(re.findall(r"\b[\w%-]+\b", paper["summary"]))
    assert 100 <= count <= 120, f"Summary word count {count}: {paper['citation']}"
    print(f"{paper['citation']}: {count} words")


document = Document()
section = document.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.55)
section.right_margin = Inches(0.55)
section.header_distance = Inches(0.25)
section.footer_distance = Inches(0.25)

styles = document.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(9)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(2)
normal.paragraph_format.line_spacing = 1.0
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for name, size, color, before, after in [
    ("Heading 1", 12.5, BLUE, 6, 2),
    ("Heading 2", 10.5, BLUE, 4, 1.5),
]:
    style = styles[name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = color
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

paper_heading = styles.add_style("Paper Heading", WD_STYLE_TYPE.PARAGRAPH)
paper_heading.base_style = normal
paper_heading.paragraph_format.space_before = Pt(2.5)
paper_heading.paragraph_format.space_after = Pt(0.5)
paper_heading.paragraph_format.keep_with_next = True

paper_summary = styles.add_style("Paper Summary", WD_STYLE_TYPE.PARAGRAPH)
paper_summary.base_style = normal
paper_summary.paragraph_format.space_after = Pt(1.5)
paper_summary.paragraph_format.line_spacing = 0.95
paper_summary.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

reference_style = styles.add_style("Compact Reference", WD_STYLE_TYPE.PARAGRAPH)
reference_style.base_style = normal
reference_style.paragraph_format.left_indent = Inches(0.14)
reference_style.paragraph_format.first_line_indent = Inches(-0.14)
reference_style.paragraph_format.space_after = Pt(0.5)
reference_style.paragraph_format.line_spacing = 0.9

header_p = section.header.paragraphs[0]
header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
header_p.paragraph_format.space_after = Pt(0)
header_run = header_p.add_run("CSE427  |  Project Update 1  |  Lab Sections 07/08")
set_font(header_run, size=7.5, bold=True, color=MUTED)

footer_p = section.footer.paragraphs[0]
add_page_number(footer_p)

title = document.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(1)
title_run = title.add_run("ECG Arrhythmia Classification Using CNN-LSTM")
set_font(title_run, size=16, bold=True, color=INK)

subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(4)
subtitle_run = subtitle.add_run("Project Update 1  |  10 August 2026")
set_font(subtitle_run, size=8.5, bold=True, color=MUTED)

document.add_heading("1. Introduction", level=1)

p = document.add_paragraph()
r = p.add_run("Problem statement. ")
set_font(r, size=9, bold=True, color=INK)
r = p.add_run(
    "Manual analysis of long ambulatory electrocardiograms is time-consuming, and subtle beat morphology "
    "or rhythm changes can be difficult to classify consistently. This project addresses automatic "
    "beat-level arrhythmia classification using a hybrid CNN-LSTM model. The CNN will learn local waveform "
    "features, while the LSTM will model temporal dependencies. The system is a research classifier, not a "
    "clinical diagnostic device, and will be evaluated with leakage-aware record separation and class-sensitive metrics."
)
set_font(r, size=9)

p = document.add_paragraph()
r = p.add_run("Project objectives. ")
set_font(r, size=9, bold=True, color=INK)
r = p.add_run(
    "The objectives are to (1) construct a reproducible MIT-BIH loading and preprocessing pipeline; "
    "(2) examine label imbalance, leads, and signal quality; (3) compare CNN, LSTM, and CNN-LSTM models; "
    "(4) prevent patient-specific leakage through record-wise evaluation; and (5) report accuracy together "
    "with macro-F1, per-class sensitivity, specificity, precision, and confusion matrices."
)
set_font(r, size=9)

document.add_heading("2. Short Literature Review", level=1)
intro = document.add_paragraph()
intro.paragraph_format.space_after = Pt(1.5)
intro_run = intro.add_run(
    "Eight papers are divided into two sets of four for the two group members. Each review summarizes the "
    "method, reported performance, limitation, and research gap relevant to this project."
)
set_font(intro_run, size=8.8, italic=True, color=MUTED)

document.add_heading("Paper set 1 (papers 1-4)", level=2)
for paper in PAPERS[:4]:
    add_paper_review(document, paper)

document.add_page_break()
document.add_heading("Paper set 2 (papers 5-8)", level=2)
for paper in PAPERS[4:]:
    add_paper_review(document, paper)

document.add_heading("3. Dataset Identified", level=1)
p = document.add_paragraph()
r = p.add_run("MIT-BIH Arrhythmia Database v1.0.0. ")
set_font(r, size=9, bold=True, color=INK)
r = p.add_run(
    "The open-access PhysioNet database contains 48 half-hour excerpts of two-channel ambulatory ECG from "
    "47 subjects. Signals were digitized at 360 samples per second per channel with 11-bit resolution over a "
    "10 mV range. Approximately 110,000 beat annotations were independently prepared by cardiologists. Data "
    "are distributed in WFDB format: .hea headers, .dat waveforms, and .atr reference annotations. The dataset "
    "is suitable because it is expert-labeled, manageable in Colab, widely benchmarked, and includes both common "
    "and clinically significant arrhythmias. Initial EDA audits all records and raw symbols; preprocessing uses "
    "provisional AAMI-style groups and one-second annotation-centered windows."
)
set_font(r, size=9)

source_p = document.add_paragraph()
source_p.paragraph_format.space_after = Pt(2)
source_label = source_p.add_run("Source: ")
set_font(source_label, size=8.5, bold=True)
add_hyperlink(source_p, "PhysioNet MIT-BIH Arrhythmia Database", "https://physionet.org/content/mitdb/1.0.0/")
doi_run = source_p.add_run("  |  DOI: 10.13026/C2F305  |  License: Open Data Commons Attribution 1.0")
set_font(doi_run, size=8.5)

document.add_heading("References", level=1)
for index, reference in enumerate(REFERENCES, start=1):
    p = document.add_paragraph(style="Compact Reference")
    r = p.add_run(f"[{index}] {reference}")
    set_font(r, size=7.3)

document.core_properties.title = "ECG Arrhythmia Classification Using CNN-LSTM - Project Update 1"
document.core_properties.subject = "CSE427 Week 4 project report"
document.core_properties.author = "CSE427 Project Group"
document.core_properties.keywords = "ECG, MIT-BIH, CNN-LSTM, arrhythmia classification"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
print(f"Wrote {OUTPUT}")
